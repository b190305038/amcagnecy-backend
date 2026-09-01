"""
MTD / YTD product-wise sales report — one row per Dealer, with a Volume +
Value column pair per product, grouped under "MTD Sales" / "YTD Sales"
headings. Backs both the on-screen report table (JSON) and its Excel export
(multi-level merged header, matching the client-supplied template).

New products need no code change: the product list (and therefore the
column set) is read from the Product table at request time.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Area, Dealer, InvoiceLine, Product, Shop, Territory

BASE_HEADERS = [
    "SL", "Region", "Area", "Area Code", "Territory", "Territory Code",
    "ASE Name", "ASE Code", "Dealer Code", "Dealer Name", "Contact Number", "Address",
]

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
GROUP_FILL = PatternFill("solid", fgColor="2E5B8A")
PRODUCT_FILL = PatternFill("solid", fgColor="3E76AC")


def _period_bounds(as_of: date) -> tuple[str, str, str]:
    """(year_start_iso, month_start_iso, as_of_iso) — dates are stored as ISO
    strings, so lexicographic comparison is chronological comparison."""
    return as_of.replace(month=1, day=1).isoformat(), as_of.replace(day=1).isoformat(), as_of.isoformat()


def volume_text(product: Product, qty: float) -> str:
    qty = int(round(qty))
    if qty <= 0:
        return "0"
    if (product.unit_type or "PCS") != "Carton":
        return f"{qty} PCS"
    per = max(1, int(product.pcs_per_carton or 1))
    cartons, extra = divmod(qty, per)
    if cartons and extra:
        return f"{cartons} Carton + {extra} PCS"
    if cartons:
        return f"{cartons} Carton"
    return f"{extra} PCS"


def _area_codes() -> dict:
    return {a.name: a.code for a in Area.objects.all()}


def _territory_codes() -> dict:
    return {(t.area.name, t.name): t.code for t in Territory.objects.select_related("area").all()}


def build_report(as_of: date, *, region: str = "", area: str = "", territory: str = "") -> dict:
    """Everything the JSON view and the Excel export both need."""
    year_start, month_start, as_of_iso = _period_bounds(as_of)

    products = list(Product.objects.all().order_by("brand", "name"))
    dealers = Dealer.objects.select_related("ase").order_by("region", "area", "territory", "name")
    if region:
        dealers = dealers.filter(region=region)
    if area:
        dealers = dealers.filter(area=area)
    if territory:
        dealers = dealers.filter(territory=territory)
    dealers = list(dealers)

    # shop *name* -> dealer code. SalesOrder.shop stores the shop's name, not
    # its code, so the join has to happen on that name in Python — there's no
    # FK from SalesOrder to Shop to walk.
    shop_dealer = dict(Shop.objects.exclude(dealer=None).values_list("name", "dealer_id"))

    lines = InvoiceLine.objects.filter(
        order__date__gte=year_start, order__date__lte=as_of_iso
    ).values("sku", "qty", "price", "order__shop", "order__date")

    # dealer_code -> sku -> {mtd_qty, mtd_val, ytd_qty, ytd_val}
    agg: dict = defaultdict(lambda: defaultdict(lambda: {"mtd_qty": 0.0, "mtd_val": 0.0, "ytd_qty": 0.0, "ytd_val": 0.0}))
    for row in lines:
        dealer_code = shop_dealer.get(row["order__shop"])
        if not dealer_code:
            continue
        qty = float(row["qty"] or 0)
        value = qty * float(row["price"] or 0)
        bucket = agg[dealer_code][row["sku"]]
        bucket["ytd_qty"] += qty
        bucket["ytd_val"] += value
        if row["order__date"] >= month_start:
            bucket["mtd_qty"] += qty
            bucket["mtd_val"] += value

    area_codes = _area_codes()
    territory_codes = _territory_codes()

    rows = []
    for idx, dealer in enumerate(dealers, start=1):
        dealer_agg = agg.get(dealer.code, {})
        mtd, ytd = {}, {}
        for product in products:
            bucket = dealer_agg.get(product.sku, {"mtd_qty": 0, "mtd_val": 0, "ytd_qty": 0, "ytd_val": 0})
            mtd[product.sku] = {
                "volume": volume_text(product, bucket["mtd_qty"]),
                "value": round(bucket["mtd_val"], 2),
            }
            ytd[product.sku] = {
                "volume": volume_text(product, bucket["ytd_qty"]),
                "value": round(bucket["ytd_val"], 2),
            }
        rows.append({
            "sl": idx,
            "region": dealer.region,
            "area": dealer.area,
            "areaCode": area_codes.get(dealer.area, ""),
            "territory": dealer.territory,
            "territoryCode": territory_codes.get((dealer.area, dealer.territory), ""),
            "aseName": dealer.ase.name if dealer.ase else "",
            "aseCode": dealer.ase_id or "",
            "dealerCode": dealer.code,
            "dealerName": dealer.name,
            "contactNumber": dealer.contact_number,
            "address": dealer.address,
            "mtd": mtd,
            "ytd": ytd,
        })

    return {
        "asOf": as_of_iso,
        "monthStart": month_start,
        "yearStart": year_start,
        "products": [{"sku": p.sku, "label": f"{p.brand} {p.name}".strip()} for p in products],
        "rows": rows,
    }


# ---------------------------------------------------------------------------
# Excel export: 3-row merged header (basic info | MTD Sales | YTD Sales, each
# product spanning a Volume + Value pair), then one row per dealer.
# ---------------------------------------------------------------------------
def build_report_xlsx(report: dict) -> bytes:
    products = report["products"]
    wb = Workbook()
    ws = wb.active
    ws.title = "MTD & YTD Sales"

    n_base = len(BASE_HEADERS)
    n_products = len(products)
    period_span = n_products * 2  # Volume + Value per product

    # Row 1: basic headers (merged vertically across the 3 header rows) + the
    # two period group headings (merged horizontally across their columns).
    for idx, header in enumerate(BASE_HEADERS, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=1, start_column=idx, end_row=3, end_column=idx)

    for period, label in (("MTD", "MTD Sales"), ("YTD", "YTD Sales")):
        start_col = n_base + 1 + (0 if period == "MTD" else period_span)
        end_col = start_col + period_span - 1
        cell = ws.cell(row=1, column=start_col, value=label)
        cell.font = Font(bold=True, color="FFFFFF", size=12)
        cell.fill = GROUP_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if period_span > 1:
            ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)

        # Row 2: one merged cell per product (spans its Volume + Value pair).
        for p_idx, product in enumerate(products):
            col = start_col + p_idx * 2
            cell = ws.cell(row=2, column=col, value=product["label"])
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PRODUCT_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 1)

            # Row 3: Volume / Value sub-headers.
            for sub_idx, sub in enumerate(("Volume", "Value")):
                cell = ws.cell(row=3, column=col + sub_idx, value=sub)
                cell.font = Font(bold=True, color="FFFFFF", size=10)
                cell.fill = PRODUCT_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 18
    ws.freeze_panes = "A4"

    ws.column_dimensions["A"].width = 6
    for idx, header in enumerate(BASE_HEADERS[1:], start=2):
        ws.column_dimensions[get_column_letter(idx)].width = max(12, len(header) + 2)
    for col in range(n_base + 1, n_base + 1 + 2 * period_span):
        ws.column_dimensions[get_column_letter(col)].width = 16

    row_at = 4
    for row in report["rows"]:
        values = [
            row["sl"], row["region"], row["area"], row["areaCode"], row["territory"],
            row["territoryCode"], row["aseName"], row["aseCode"], row["dealerCode"],
            row["dealerName"], row["contactNumber"], row["address"],
        ]
        for period_key in ("mtd", "ytd"):
            for product in products:
                cell = row[period_key][product["sku"]]
                values.append(cell["volume"])
                values.append(cell["value"])
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_at, column=col, value=value)
        row_at += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
